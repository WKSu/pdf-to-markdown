# /// script
# requires-python = ">=3.12"
# dependencies = [
#     "pymupdf==1.27.2.2",
#     "pymupdf4llm==0.3.4",
#     "tabulate",
#     "python-docx",
#     "python-pptx",
# ]
# ///
"""Build Word and PowerPoint test documents and check what comes out.

The counterpart of make-sample.py for the Office path. Each document carries
the cases that a naive extraction gets wrong -- footnotes, tracked changes,
merged cells, list styles that nest by indentation, columns stored right
before left, a chart, speaker notes, a hidden and a reordered slide -- and the
checks below say what the Markdown must contain. python-docx and python-pptx
only build the files; the converter itself uses the standard library.

    uv run tests/office-check.py

Writes test-pdfs/sample.docx and test-pdfs/sample.pptx, which browser-test.mjs
can then take as input.
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT / "python"))

import pymupdf  # noqa: E402
from convert import Options, assemble  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402
from docx.shared import Inches  # noqa: E402
from office import OfficeConverter, detect_format  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches as PptInches  # noqa: E402

OUT = ROOT / "test-pdfs"
HYPERLINK = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)


def png() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 40), 0)
    pix.set_rect(pix.irect, (0, 130, 70))
    return pix.tobytes("png")


# --------------------------------------------------------------------- Word


def build_docx() -> bytes:
    d = Document()
    d.add_paragraph("Concept, niet verspreiden.")
    d.add_heading("Aanleiding", 1)
    p = d.add_paragraph("De stad groeit. ")
    p.add_run("Dit is vet").bold = True
    p.add_run(" en dit is ")
    p.add_run("cursief").italic = True
    p.add_run(". Zie ")
    rid = d.part.relate_to("https://www.rotterdam.nl", HYPERLINK, is_external=True)
    p._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}">'
            "<w:r><w:t>rotterdam.nl</w:t></w:r></w:hyperlink>"
        )
    )
    p.add_run(" en de voetnoot.")
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:footnoteReference w:id="1"/></w:r>'))
    p._p.append(
        parse_xml(
            f'<w:del {nsdecls("w")} w:id="9" w:author="x">'
            "<w:r><w:delText>VERWIJDERDE TEKST</w:delText></w:r></w:del>"
        )
    )
    d.add_paragraph("2024. Een jaar, geen lijst.")

    d.add_heading("Doelen", 1)
    d.add_paragraph("Veilig", style="List Number")
    d.add_paragraph("Fiets", style="List Number 2")
    d.add_paragraph("Bereikbaar", style="List Number")
    d.add_paragraph("Punt A", style="List Bullet")
    d.add_paragraph("Sub A1", style="List Bullet 2")
    t = d.add_table(rows=2, cols=3)
    t.style = "Table Grid"
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Modal split"
    t.cell(0, 2).text = "2024"
    t.cell(1, 0).text = "Fiets"
    t.cell(1, 1).text = "auto | ov"
    t.cell(1, 2).text = "30%"
    image = OUT / "_tmp.png"
    image.write_bytes(png())
    d.add_picture(str(image), width=Inches(1))
    image.unlink()
    d.inline_shapes[0]._inline.docPr.set("descr", "Kaart van de fietsroutes")
    d.add_heading("Ambitie", 1)
    d.add_paragraph("Tweede voetnoot volgt.").runs[0]._r.addnext(
        parse_xml(f'<w:r {nsdecls("w")}><w:footnoteReference w:id="2"/></w:r>')
    )

    buffer = io.BytesIO()
    d.save(buffer)
    return add_footnotes(
        buffer.getvalue(),
        {"1": "Bron: Mobiliteitsmonitor 2025.", "2": "Tweede noot."},
    )


def add_footnotes(docx: bytes, notes: dict[str, str]) -> bytes:
    """python-docx cannot write footnotes, so add the part by hand."""
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    body = "".join(
        f'<w:footnote w:id="{i}"><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:footnote>'
        for i, text in notes.items()
    )
    footnotes = f'<?xml version="1.0"?><w:footnotes {w}>{body}</w:footnotes>'
    src = zipfile.ZipFile(io.BytesIO(docx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/word/footnotes.xml" ContentType="application/'
                    b'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
                    b"</Types>",
                )
            if item.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
                    b"</Relationships>",
                )
            dst.writestr(item, data)
        dst.writestr("word/footnotes.xml", footnotes)
    return out.getvalue()


# --------------------------------------------------------------- PowerPoint


def build_pptx() -> bytes:
    p = Presentation()
    blank = p.slide_layouts[6]

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "Doelen fietsbeleid"
    tf = s.placeholders[1].text_frame
    tf.text = "Veiligheid"
    for text, level in [("Minder ongevallen", 1), ("Bereikbaarheid", 0)]:
        para = tf.add_paragraph()
        para.text = text
        para.level = level
    s.notes_slide.notes_text_frame.text = "Benadruk de stijging in 2025."

    s = p.slides.add_slide(blank)
    # Written right column first: the XML order is not the reading order.
    for x, text in [(6, "Rechts: conclusie"), (0.5, "Links: aanleiding")]:
        box = s.shapes.add_textbox(
            PptInches(x), PptInches(2), PptInches(3), PptInches(1)
        )
        box.text_frame.text = text
    title = s.shapes.add_textbox(
        PptInches(0.5), PptInches(0.3), PptInches(8), PptInches(0.8)
    )
    title.text_frame.text = "Ontwikkeling fietsgebruik"
    data = CategoryChartData()
    data.categories = ["2022", "2023", "2024"]
    data.add_series("Fietsritten (mln)", (41.2, 43.5, 46.1))
    s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        PptInches(0.5),
        PptInches(3.5),
        PptInches(5),
        PptInches(3),
        data,
    )
    picture = s.shapes.add_picture(io.BytesIO(png()), PptInches(6), PptInches(4))
    picture._element.nvPicPr.cNvPr.set("descr", "Foto van een fietsstraat")

    s = p.slides.add_slide(blank)
    s.shapes.add_textbox(
        PptInches(0.5), PptInches(0.3), PptInches(8), PptInches(0.8)
    ).text_frame.text = "Tabel"
    table = s.shapes.add_table(
        2, 2, PptInches(0.5), PptInches(1.5), PptInches(4), PptInches(1)
    ).table
    for (r, c), text in {
        (0, 0): "Gebied",
        (0, 1): "Fiets",
        (1, 0): "Noord",
        (1, 1): "34%",
    }.items():
        table.cell(r, c).text = text
    s._element.set("show", "0")  # hidden slide

    # Move the last slide to the front: order comes from sldIdLst, not names.
    ids = p.slides._sldIdLst
    last = list(ids)[-1]
    ids.remove(last)
    ids.insert(0, last)

    buffer = io.BytesIO()
    p.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------- run


def convert(data: bytes, name: str, **opts) -> str:
    conv = OfficeConverter(data, name, Options(**opts))
    pages = [conv.page(n) for n in range(1, conv.page_count + 1)]
    return assemble(conv.front_matter(), pages, True, conv.unit)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    failures = []

    def expect(label: str, ok: bool) -> None:
        print(f"  {'ok  ' if ok else 'FAIL'} {label}")
        if not ok:
            failures.append(label)

    docx = build_docx()
    (OUT / "sample.docx").write_bytes(docx)
    md = convert(docx, "sample.docx")
    print("sample.docx")
    expect("detected as docx", detect_format(docx) == "docx")
    expect("split into 4 sections", md.count("<!-- section ") == 4)
    expect("headings", "# Aanleiding" in md and "# Doelen" in md)
    expect("bold and italic", "**Dit is vet**" in md and "*cursief*" in md)
    expect("hyperlink", "[rotterdam.nl](https://www.rotterdam.nl)" in md)
    expect("tracked deletion dropped", "VERWIJDERDE" not in md)
    expect("number-dot text escaped", "2024\\. Een jaar" in md)
    expect("nested numbered list", "1. Veilig\n   1. Fiets\n2. Bereikbaar" in md)
    expect("nested bullets", "- Punt A\n   - Sub A1" in md)
    expect("merged cell padded", "| Modal split |  | 2024 |" in md)
    expect("pipe in cell escaped", "auto \\| ov" in md)
    expect(
        "image with alt text",
        "![Kaart van de fietsroutes](figures/sample-img1.png)" in md,
    )
    expect(
        "footnote ref and definition",
        "[^1]" in md and "[^1]: Bron: Mobiliteitsmonitor 2025." in md,
    )
    sections = md.split("<!-- section ")
    expect(
        "footnote 2 defined in its own section", "[^2]: Tweede noot." in sections[-1]
    )

    pptx = build_pptx()
    (OUT / "sample.pptx").write_bytes(pptx)
    md = convert(pptx, "sample.pptx")
    print("sample.pptx")
    expect("detected as pptx", detect_format(pptx) == "pptx")
    expect("3 slides", md.count("<!-- slide ") == 3)
    expect(
        "reordered slide comes first",
        md.index("## Dia 1: Tabel") < md.index("## Dia 2: Doelen"),
    )
    expect("hidden slide marked", "## Dia 1: Tabel _(verborgen dia)_" in md)
    expect(
        "nested bullets from placeholder",
        "- Veiligheid\n   - Minder ongevallen\n- Bereikbaarheid" in md,
    )
    expect("speaker notes", "> Benadruk de stijging in 2025." in md)
    expect("title from text box", "## Dia 3: Ontwikkeling fietsgebruik" in md)
    expect(
        "left column before right",
        md.index("Links: aanleiding") < md.index("Rechts: conclusie"),
    )
    expect("chart as table", "| 2024 | 46.1 |" in md)
    expect("table", "| Noord | 34% |" in md)
    expect(
        "picture with alt text",
        "![Foto van een fietsstraat](figures/sample-dia3-img1.png)" in md,
    )

    quiet = convert(pptx, "sample.pptx", speaker_notes=False, ignore_images=True)
    expect(
        "notes can be switched off",
        "Benadruk" not in quiet and "speaker_notes: false" in quiet,
    )
    expect("images can be switched off", "![" not in quiet)

    print("detection")
    try:
        detect_format(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(100), "oud.doc")
        expect("legacy .doc refused", False)
    except ValueError as error:
        expect("legacy .doc refused", "Oud Office-formaat" in str(error))

    print(f"\n{'FAIL' if failures else 'PASS'}: {len(failures)} failed")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
